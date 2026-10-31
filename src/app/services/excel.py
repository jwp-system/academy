"""엑셀 입출력: 학생 명단 가져오기(2주차), 수납 내역 내보내기(5주차), 거래내역 업로드(7주차)."""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from io import BytesIO
from typing import BinaryIO

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlmodel import Session, select

from app.models import Guardian, Student, StudentStatus
from app.schemas import GuardianIn, StudentIn, form_errors
from app.services.students import create_student

# 엑셀 머리글 → (스키마, 필드명). 열 순서는 상관없다.
STUDENT_COLUMNS: dict[str, tuple[str, str]] = {
    "이름": ("student", "name"),
    "학교": ("student", "school"),
    "학년": ("student", "grade"),
    "학생연락처": ("student", "phone"),
    "상태": ("student", "status"),
    "메모": ("student", "memo"),
    "보호자명": ("guardian", "name"),
    "보호자관계": ("guardian", "relation"),
    "보호자연락처": ("guardian", "phone"),
    "보호자이메일": ("guardian", "email"),
}
STATUS_LABELS = {
    "재원": StudentStatus.active,
    "휴원": StudentStatus.paused,
    "퇴원": StudentStatus.withdrawn,
}
FIELD_LABELS = {(kind, name): header for header, (kind, name) in STUDENT_COLUMNS.items()}


@dataclass
class ImportRow:
    row_no: int
    student: StudentIn | None = None
    guardian: GuardianIn | None = None
    errors: list[str] = field(default_factory=list)
    duplicate: bool = False


@dataclass
class ImportResult:
    rows: list[ImportRow]
    created: int = 0
    header_error: str | None = None

    @property
    def errors(self) -> list[ImportRow]:
        return [r for r in self.rows if r.errors]

    @property
    def duplicates(self) -> list[ImportRow]:
        return [r for r in self.rows if r.duplicate]


def student_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "학생"
    ws.append(list(STUDENT_COLUMNS))
    ws.append(
        ["홍길동", "OO중", "중2", "010-1234-5678", "재원", "", "홍판서", "부", "010-9876-5432", ""]
    )
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 14
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _validate(kind: str, schema, data: dict, row: ImportRow):
    try:
        return schema.model_validate(data)
    except ValidationError as e:
        for name, msg in form_errors(e).items():
            row.errors.append(f"{FIELD_LABELS.get((kind, name), name)}: {msg}")
        return None


def parse_students(file: BinaryIO) -> ImportResult:
    try:
        ws = load_workbook(file, read_only=True, data_only=True).worksheets[0]
    except Exception:
        return ImportResult(rows=[], header_error="엑셀(.xlsx) 파일을 읽을 수 없어요.")
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows_iter, ())]
    if "이름" not in header:
        return ImportResult(
            rows=[], header_error="첫 줄에 '이름' 머리글이 필요해요. 양식을 확인해 주세요."
        )
    columns = {i: STUDENT_COLUMNS[h] for i, h in enumerate(header) if h in STUDENT_COLUMNS}

    rows: list[ImportRow] = []
    for row_no, values in enumerate(rows_iter, start=2):
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        data: dict[str, dict] = {"student": {}, "guardian": {}}
        for i, (kind, name) in columns.items():
            value = values[i] if i < len(values) else None
            data[kind][name] = str(value) if isinstance(value, int) and name != "phone" else value
        status = data["student"].get("status")
        if isinstance(status, str) and status.strip() in STATUS_LABELS:
            data["student"]["status"] = STATUS_LABELS[status.strip()]
        elif not status:
            data["student"].pop("status", None)

        row = ImportRow(row_no=row_no)
        row.student = _validate("student", StudentIn, data["student"], row)
        if any(v not in (None, "") for v in data["guardian"].values()):
            row.guardian = _validate(
                "guardian", GuardianIn, data["guardian"] | {"notify_sms": True}, row
            )
        rows.append(row)
    return ImportResult(rows=rows)


def import_students(db: Session, file: BinaryIO) -> ImportResult:
    """오류가 한 줄이라도 있으면 아무것도 등록하지 않는다. 이름+연락처가 같은 기존 학생은 건너뛴다."""
    result = parse_students(file)
    if result.header_error or result.errors:
        return result
    existing_keys: set[tuple[str, str]] = set()
    for name, phone, gphone in db.exec(
        select(Student.name, Student.phone, Guardian.phone).join(
            Guardian, Guardian.student_id == Student.id, isouter=True
        )
    ).all():
        existing_keys |= {(name, p) for p in (phone, gphone) if p}
    for row in result.rows:
        key_phone = row.student.phone or (row.guardian.phone if row.guardian else None)
        if key_phone and (row.student.name, key_phone) in existing_keys:
            row.duplicate = True
            continue
        create_student(db, row.student, [row.guardian] if row.guardian else [])
        if key_phone:
            existing_keys.add((row.student.name, key_phone))
        result.created += 1
    db.commit()
    return result


# --- 수납 내역 내보내기 (5주차) ------------------------------------------------------


def _style_sheet(ws, widths: list[int], money_cols: list[int]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F2F4F6")
        cell.alignment = Alignment(horizontal="center")
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for i in money_cols:
            row[i - 1].number_format = "#,##0"
    ws.freeze_panes = "A2"


def billing_workbook(month_label: str, rows, payments, labels: dict, methods: dict) -> bytes:
    """rows: services.billing.BillingRow 목록, payments: (Payment, Invoice, Student) 목록."""
    wb = Workbook()
    ws = wb.active
    ws.title = "청구"
    ws.append(
        [
            "학생",
            "학교",
            "수강 반",
            "입금자명",
            "수강료",
            "할인",
            "청구액",
            "납부액",
            "남은 금액",
            "납부 기한",
            "상태",
            "최근 납부일",
            "결제 수단",
        ]
    )
    for r in rows:
        ws.append(
            [
                r.student.name,
                r.student.school or "",
                ", ".join(c.name for c in r.classes),
                r.student.payer_code,
                r.invoice.amount,
                r.invoice.discount,
                r.due,
                r.paid,
                r.rest,
                r.invoice.due_date,
                labels[r.invoice.status],
                r.last_payment.paid_at.date() if r.last_payment else None,
                methods[r.last_payment.method] if r.last_payment else "",
            ]
        )
    n = len(rows) + 1
    ws.append([])
    ws.append(
        [
            "합계",
            "",
            "",
            "",
            f"=SUM(E2:E{n})",
            f"=SUM(F2:F{n})",
            f"=SUM(G2:G{n})",
            f"=SUM(H2:H{n})",
            f"=SUM(I2:I{n})",
        ]
    )
    _style_sheet(ws, [10, 12, 26, 12, 11, 9, 11, 11, 11, 12, 10, 12, 10], [5, 6, 7, 8, 9])
    for cell in ws[ws.max_row]:
        cell.number_format = "#,##0"
    for c in ws.iter_rows(min_row=2, min_col=10, max_col=10):
        c[0].number_format = "yyyy-mm-dd"
    for c in ws.iter_rows(min_row=2, min_col=12, max_col=12):
        c[0].number_format = "yyyy-mm-dd"

    ws2 = wb.create_sheet("납부 내역")
    ws2.append(["납부일시", "학생", "청구월", "금액", "결제 수단", "현금영수증", "메모"])
    for p, inv, s in payments:
        receipt = ("발급" if p.cash_receipt_issued else "미발급") if p.method != "card" else "-"
        ws2.append(
            [p.paid_at, s.name, inv.year_month, p.amount, methods[p.method], receipt, p.memo or ""]
        )
    _style_sheet(ws2, [17, 10, 10, 11, 10, 11, 30], [4])
    for c in ws2.iter_rows(min_row=2, max_col=1):
        c[0].number_format = "yyyy-mm-dd hh:mm"
    ws.title = f"{month_label} 청구"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- 거래내역 업로드 (7주차) ------------------------------------------------------

# 은행마다 머리글이 달라서 키워드로 열을 찾는다. 앞쪽 키워드가 우선.
_STATEMENT_COLUMNS = {
    "datetime": ("거래일시", "거래일자", "거래일", "일시", "날짜"),
    "time": ("거래시간", "시간", "시각"),
    "deposit": ("입금액", "입금금액", "맡기신금액", "입금(원)", "입금"),
    "name": ("기재내용", "입금자명", "보낸분", "의뢰인", "내용", "적요", "거래내용", "메모"),
}


@dataclass
class StatementRow:
    row_no: int
    received_at: datetime
    payer_name: str
    amount: int


@dataclass
class StatementResult:
    rows: list[StatementRow] = field(default_factory=list)
    error: str | None = None


def _read_rows(data: bytes, filename: str) -> list[list]:
    name = filename.lower()
    if name.endswith(".csv"):
        for enc in ("utf-8-sig", "cp949"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("CSV 파일의 글자 인코딩을 알 수 없어요.")
        return [row for row in csv.reader(io.StringIO(text))]
    if name.endswith(".xls"):
        raise ValueError(
            "예전 엑셀(.xls) 파일이에요. 엑셀에서 '다른 이름으로 저장 → .xlsx'로 저장한 뒤 올려 주세요."
        )
    try:
        ws = load_workbook(BytesIO(data), read_only=True, data_only=True).worksheets[0]
    except Exception:
        raise ValueError("엑셀(.xlsx)이나 CSV 파일을 읽을 수 없어요.") from None
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _find_header(rows: list[list]) -> tuple[int, dict[str, int]] | None:
    for i, row in enumerate(rows[:15]):
        cells = [str(c).replace(" ", "").strip() if c is not None else "" for c in row]
        cols: dict[str, int] = {}
        for key, words in _STATEMENT_COLUMNS.items():
            for w in words:
                idx = next((j for j, c in enumerate(cells) if c == w), None)
                if idx is None:
                    idx = next((j for j, c in enumerate(cells) if w in c and "출금" not in c), None)
                if idx is not None and idx not in cols.values():
                    cols[key] = idx
                    break
        if {"datetime", "deposit", "name"} <= cols.keys():
            return i, cols
    return None


def _parse_dt(value, time_value=None) -> datetime | None:
    if isinstance(value, datetime):
        dt_ = value
    elif isinstance(value, date):
        dt_ = datetime.combine(value, datetime.min.time())
    else:
        text = re.sub(r"\s+", " ", str(value or "").strip()).replace(".", "-").replace("/", "-")
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y%m%d %H%M%S", "%Y%m%d"):
            try:
                dt_ = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        else:
            return None
    if time_value not in (None, ""):
        if isinstance(time_value, time):
            t = time_value
        else:
            m = re.match(r"(\d{1,2}):?(\d{2})(?::?(\d{2}))?", str(time_value).strip())
            t = time(int(m.group(1)), int(m.group(2)), int(m.group(3) or 0)) if m else None
        if t:
            dt_ = datetime.combine(dt_.date(), t)
    return dt_


def _money(value) -> int:
    if isinstance(value, (int, float)):
        return int(value)
    digits = re.sub(r"[^\d]", "", str(value or ""))
    return int(digits) if digits else 0


def _cell(row: list, cols: dict[str, int], key: str):
    j = cols.get(key)
    return row[j] if j is not None and j < len(row) else None


def parse_bank_statement(data: bytes, filename: str) -> StatementResult:
    """입금 행만 뽑는다 (출금·입금 0원 행은 건너뜀)."""
    try:
        rows = _read_rows(data, filename)
    except ValueError as e:
        return StatementResult(error=str(e))
    found = _find_header(rows)
    if not found:
        return StatementResult(
            error="거래일시·입금액·입금자(기재내용) 열을 찾지 못했어요. 인터넷뱅킹에서 받은 거래내역 파일인지 확인해 주세요."
        )
    header_idx, cols = found
    out = StatementResult()
    for i, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        amount = _money(_cell(row, cols, "deposit"))
        name = str(_cell(row, cols, "name") or "").strip()
        when = _parse_dt(_cell(row, cols, "datetime"), _cell(row, cols, "time"))
        if amount <= 0 or not name or when is None:
            continue
        out.rows.append(StatementRow(i, when, name, amount))
    return out
