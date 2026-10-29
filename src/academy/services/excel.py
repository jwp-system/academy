"""엑셀 입출력: 학생 명단 가져오기(2주차), 수납 내역 내보내기(5주차), 거래내역 업로드(7주차)."""

from dataclasses import dataclass, field
from io import BytesIO
from typing import BinaryIO

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlmodel import Session, select

from academy.models import Guardian, Student, StudentStatus
from academy.schemas import GuardianIn, StudentIn, form_errors
from academy.services.students import create_student

# 엑셀 머리글 → (스키마, 필드명). 열 순서는 상관없다.
STUDENT_COLUMNS: dict[str, tuple[str, str]] = {
    "이름": ("student", "name"),
    "학교": ("student", "school"),
    "학년": ("student", "grade"),
    "학생연락처": ("student", "phone"),
    "상태": ("student", "status"),
    "메모": ("student", "memo"),
    "보호자명": ("guardian", "name"),
    "보호자관계": ("guardian", "relation"),
    "보호자연락처": ("guardian", "phone"),
    "보호자이메일": ("guardian", "email"),
}
STATUS_LABELS = {
    "재원": StudentStatus.active,
    "휴원": StudentStatus.paused,
    "퇴원": StudentStatus.withdrawn,
}
FIELD_LABELS = {(kind, name): header for header, (kind, name) in STUDENT_COLUMNS.items()}


@dataclass
class ImportRow:
    row_no: int
    student: StudentIn | None = None
    guardian: GuardianIn | None = None
    errors: list[str] = field(default_factory=list)
    duplicate: bool = False


@dataclass
class ImportResult:
    rows: list[ImportRow]
    created: int = 0
    header_error: str | None = None

    @property
    def errors(self) -> list[ImportRow]:
        return [r for r in self.rows if r.errors]

    @property
    def duplicates(self) -> list[ImportRow]:
        return [r for r in self.rows if r.duplicate]


def student_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "학생"
    ws.append(list(STUDENT_COLUMNS))
    ws.append(
        ["홍길동", "OO중", "중2", "010-1234-5678", "재원", "", "홍판서", "부", "010-9876-5432", ""]
    )
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 14
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _validate(kind: str, schema, data: dict, row: ImportRow):
    try:
        return schema.model_validate(data)
    except ValidationError as e:
        for name, msg in form_errors(e).items():
            row.errors.append(f"{FIELD_LABELS.get((kind, name), name)}: {msg}")
        return None


def parse_students(file: BinaryIO) -> ImportResult:
    try:
        ws = load_workbook(file, read_only=True, data_only=True).worksheets[0]
    except Exception:
        return ImportResult(rows=[], header_error="엑셀(.xlsx) 파일을 읽을 수 없어요.")
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows_iter, ())]
    if "이름" not in header:
        return ImportResult(
            rows=[], header_error="첫 줄에 '이름' 머리글이 필요해요. 양식을 확인해 주세요."
        )
    columns = {i: STUDENT_COLUMNS[h] for i, h in enumerate(header) if h in STUDENT_COLUMNS}

    rows: list[ImportRow] = []
    for row_no, values in enumerate(rows_iter, start=2):
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        data: dict[str, dict] = {"student": {}, "guardian": {}}
        for i, (kind, name) in columns.items():
            value = values[i] if i < len(values) else None
            data[kind][name] = str(value) if isinstance(value, int) and name != "phone" else value
        status = data["student"].get("status")
        if isinstance(status, str) and status.strip() in STATUS_LABELS:
            data["student"]["status"] = STATUS_LABELS[status.strip()]
        elif not status:
            data["student"].pop("status", None)

        row = ImportRow(row_no=row_no)
        row.student = _validate("student", StudentIn, data["student"], row)
        if any(v not in (None, "") for v in data["guardian"].values()):
            row.guardian = _validate(
                "guardian", GuardianIn, data["guardian"] | {"notify_sms": True}, row
            )
        rows.append(row)
    return ImportResult(rows=rows)


def import_students(db: Session, file: BinaryIO) -> ImportResult:
    """오류가 한 줄이라도 있으면 아무것도 등록하지 않는다. 이름+연락처가 같은 기존 학생은 건너뛴다."""
    result = parse_students(file)
    if result.header_error or result.errors:
        return result
    existing_keys: set[tuple[str, str]] = set()
    for name, phone, gphone in db.exec(
        select(Student.name, Student.phone, Guardian.phone).join(
            Guardian, Guardian.student_id == Student.id, isouter=True
        )
    ).all():
        existing_keys |= {(name, p) for p in (phone, gphone) if p}
    for row in result.rows:
        key_phone = row.student.phone or (row.guardian.phone if row.guardian else None)
        if key_phone and (row.student.name, key_phone) in existing_keys:
            row.duplicate = True
            continue
        create_student(db, row.student, [row.guardian] if row.guardian else [])
        if key_phone:
            existing_keys.add((row.student.name, key_phone))
        result.created += 1
    db.commit()
    return result
