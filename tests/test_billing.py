import json
from datetime import date, datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlmodel import select

from academy.models import (
    BankTransaction,
    ClassGroup,
    Enrollment,
    Invoice,
    InvoiceStatus,
    Payment,
    PaymentMethod,
    StudentStatus,
)
from academy.schemas import StudentIn
from academy.services import billing as svc
from academy.services.students import create_student

NOV = date(2026, 11, 1)


def _student(db, name, classes, start=date(2026, 9, 1), end=None, status=StudentStatus.active):
    s = create_student(db, StudentIn(name=name, status=status))
    db.flush()
    for c in classes:
        db.add(Enrollment(student_id=s.id, class_id=c.id, start_date=start, end_date=end))
    db.commit()
    return s


@pytest.fixture
def classes(db):
    math = ClassGroup(name="수학", monthly_fee=280000)
    eng = ClassGroup(name="영어", monthly_fee=250000)
    free = ClassGroup(name="자습", monthly_fee=0)
    db.add_all([math, eng, free])
    db.commit()
    return math, eng, free


# --- 서비스 ------------------------------------------------------------------


def test_month_helpers():
    assert svc.parse_month("2026-02") == date(2026, 2, 1)
    assert svc.due_date_for(date(2026, 2, 1), 31) == date(2026, 2, 28)
    assert svc.shift_month(date(2026, 12, 1), 1) == date(2027, 1, 1)
    assert svc.shift_month(date(2026, 1, 1), -1) == date(2025, 12, 1)
    with pytest.raises(ValueError):
        svc.parse_month("2026-13")


def test_generate_sums_classes_and_skips(db, classes):
    math, eng, free = classes
    a = _student(db, "둘다", [math, eng])
    _student(db, "휴원", [math], status=StudentStatus.paused)
    _student(db, "지난달종료", [math], end=date(2026, 10, 31))
    _student(db, "다음달시작", [math], start=date(2026, 12, 1))
    _student(db, "중순종료", [eng], end=date(2026, 11, 15))  # 그 달에 하루라도 수강 → 청구
    _student(db, "무료반만", [free])
    assert svc.generate_invoices(db, NOV, 10) == 2
    db.commit()
    invs = {i.student_id: i for i in db.exec(select(Invoice)).all()}
    assert invs[a.id].amount == 530000 and invs[a.id].due_date == date(2026, 11, 10)
    # 다시 돌려도 중복 없음, 기존 청구는 그대로
    invs[a.id].amount = 1
    db.commit()
    assert svc.generate_invoices(db, NOV, 10) == 0
    db.refresh(invs[a.id])
    assert invs[a.id].amount == 1


def test_status_flow(db, classes):
    s = _student(db, "가", [classes[0]])
    svc.generate_invoices(db, NOV, 10)
    db.commit()
    inv = db.exec(select(Invoice)).one()
    inv.discount = 30000  # 청구액 250,000
    svc.add_payment(db, inv, amount=100000, method=PaymentMethod.cash)
    db.commit()
    assert inv.status == InvoiceStatus.partial and svc.remaining(db, inv) == 150000
    svc.mark_paid(db, inv, PaymentMethod.transfer)
    db.commit()
    assert inv.status == InvoiceStatus.paid
    assert [p.amount for p in db.exec(select(Payment).order_by(Payment.id)).all()] == [
        100000,
        150000,
    ]
    assert svc.mark_paid(db, inv, PaymentMethod.transfer) is None  # 이미 완납
    svc.clear_payments(db, inv)
    db.commit()
    assert inv.status == InvoiceStatus.unpaid and db.exec(select(Payment)).first() is None
    assert s


def test_matched_payment_locked(db, classes):
    _student(db, "가", [classes[0]])
    svc.generate_invoices(db, NOV, 10)
    inv = db.exec(select(Invoice)).one()
    p = svc.add_payment(db, inv, amount=280000, method=PaymentMethod.transfer)
    tx = BankTransaction(received_at=datetime.now(), payer_name="가1", amount=280000)
    db.add(tx)
    db.flush()
    p.bank_transaction_id = tx.id  # 입금 매칭으로 들어온 납부
    db.flush()
    with pytest.raises(svc.PaymentLocked):
        svc.clear_payments(db, inv)
    with pytest.raises(svc.PaymentLocked):
        svc.delete_payment(db, p)
    db.rollback()


def test_summary(db, classes):
    math, eng, _ = classes
    for n in ("가", "나", "다"):
        _student(db, n, [math])
    svc.generate_invoices(db, NOV, 10)
    db.commit()
    a, b, c = db.exec(select(Invoice).order_by(Invoice.id)).all()
    svc.mark_paid(db, a, PaymentMethod.card)
    svc.add_payment(db, b, amount=80000, method=PaymentMethod.cash)
    db.commit()
    rows = svc.month_rows(db, NOV)
    s = svc.summarize(db, NOV, rows, today=date(2026, 11, 20))
    assert (s.count, s.billed, s.paid, s.rest) == (3, 840000, 360000, 480000)
    assert s.unpaid_count == 2 and s.overdue_count == 2
    assert svc.summarize(db, NOV, rows, today=date(2026, 11, 5)).overdue_count == 0


# --- 화면 --------------------------------------------------------------------


@pytest.fixture
def admin(client, login):
    token = login()
    client.headers.update({"X-CSRF-Token": token, "HX-Request": "true"})
    return client


def _toast(resp):
    return json.loads(resp.headers["HX-Trigger"]).get("toast", "")


def test_generate_and_toggle_paid(admin, db, classes):
    _student(db, "홍길동", [classes[0]])
    month = svc.ym(date.today())
    page = admin.get(f"/billing?month={month}", headers={"HX-Request": ""}).text
    assert "청구를 아직 만들지 않은 재원생이 1명" in page
    resp = admin.post(f"/billing/generate?month={month}")
    assert "청구 1건을 만들었어요" in _toast(resp)
    inv = db.exec(select(Invoice)).one()

    resp = admin.patch(f"/billing/invoices/{inv.id}/paid", data={"paid": "on", "method": "cash"})
    assert "현금 완납" in _toast(resp)
    assert (
        "완납" in resp.text
        and 'id="billing-summary"' in resp.text
        and 'hx-swap-oob="true"' in resp.text
    )
    p = db.exec(select(Payment)).one()
    assert p.amount == 280000 and p.method == "cash"

    resp = admin.patch(f"/billing/invoices/{inv.id}/paid", data={"method": "cash"})  # 체크 해제
    db.refresh(inv)
    assert inv.status == "unpaid" and db.exec(select(Payment)).first() is None


def test_sheet_payment_and_invoice_edit(admin, db, classes):
    _student(db, "가", [classes[0]])
    svc.generate_invoices(db, NOV, 10)
    db.commit()
    inv = db.exec(select(Invoice)).one()
    assert "280,000" in admin.get(f"/billing/invoices/{inv.id}").text

    resp = admin.post(
        f"/billing/invoices/{inv.id}/payments", data={"amount": "0", "method": "cash"}
    )
    assert "더 큰 값을 입력해 주세요." in resp.text
    future = (date.today() + timedelta(days=1)).isoformat()
    resp = admin.post(
        f"/billing/invoices/{inv.id}/payments",
        data={"amount": "1000", "method": "cash", "paid_on": future},
    )
    assert "오늘 이후 날짜는 고를 수 없어요." in resp.text
    resp = admin.post(
        f"/billing/invoices/{inv.id}/payments",
        data={"amount": "100,000", "method": "transfer", "paid_on": "2026-10-28", "memo": "1차"},
    )
    assert "invoiceSaved" in json.loads(resp.headers["HX-Trigger"])
    db.refresh(inv)
    assert inv.status == "partial"

    resp = admin.put(
        f"/billing/invoices/{inv.id}",
        data={"amount": "280000", "discount": "300000", "due_date": "2026-11-10"},
    )
    assert "할인이 수강료보다 클 수 없어요." in resp.text
    admin.put(
        f"/billing/invoices/{inv.id}",
        data={"amount": "280000", "discount": "180000", "due_date": "2026-11-10"},
    )
    db.refresh(inv)
    assert inv.status == "paid"  # 청구액 100,000 = 이미 낸 금액

    # 납부 기록이 있으면 청구 취소 불가
    assert admin.delete(f"/billing/invoices/{inv.id}").status_code == 409
    p = db.exec(select(Payment)).one()
    admin.delete(f"/billing/payments/{p.id}")
    assert admin.delete(f"/billing/invoices/{inv.id}").status_code == 200
    assert db.exec(select(Invoice)).first() is None


def test_filters(admin, db, classes):
    for n in ("완납생", "미납생"):
        _student(db, n, [classes[0]])
    svc.generate_invoices(db, NOV, 10)
    db.commit()
    paid_inv = next(r.invoice for r in svc.month_rows(db, NOV) if r.student.name == "완납생")
    svc.mark_paid(db, paid_inv, PaymentMethod.card)
    db.commit()
    html = admin.get("/billing?month=2026-11&status=unpaid").text
    assert "미납생" in html and "완납생" not in html
    html = admin.get("/billing?month=2026-11&q=완납").text
    assert "완납생" in html and "미납생" not in html
    assert admin.get("/billing?month=2026-99").status_code == 400


def test_export(admin, db, classes):
    _student(db, "가", [classes[0], classes[1]])
    svc.generate_invoices(db, NOV, 10)
    db.commit()
    inv = db.exec(select(Invoice)).one()
    svc.add_payment(db, inv, amount=100000, method=PaymentMethod.transfer, memo="1차")
    db.commit()
    resp = admin.get("/billing/export.xlsx?month=2026-11")
    assert resp.status_code == 200 and "attachment" in resp.headers["content-disposition"]
    wb = load_workbook(BytesIO(resp.content))
    ws = wb.worksheets[0]
    assert ws.title == "11월 청구"
    row = [c.value for c in ws[2]]
    assert (
        row[0] == "가"
        and row[2] == "수학, 영어"
        and row[6] == 530000
        and row[7] == 100000
        and row[8] == 430000
    )
    assert ws.cell(row=4, column=7).value == "=SUM(G2:G2)"
    pay = [c.value for c in wb["납부 내역"][2]]
    assert pay[1] == "가" and pay[3] == 100000 and pay[5] == "미발급"


def test_receipts(admin, db, classes):
    _student(db, "가", [classes[0]])
    svc.generate_invoices(db, NOV, 10)
    inv = db.exec(select(Invoice)).one()
    t = svc.add_payment(db, inv, amount=100000, method=PaymentMethod.transfer)
    c = svc.add_payment(db, inv, amount=180000, method=PaymentMethod.card)
    db.commit()
    html = admin.get("/billing/receipts").text
    assert "100,000원" in html and "180,000원" not in html  # 카드는 대상 아님
    resp = admin.patch(f"/billing/payments/{t.id}/receipt", data={"issued": "on"})
    assert 'id="receipt-pending" hx-swap-oob="true">0<' in resp.text
    db.refresh(t)
    assert t.cash_receipt_issued
    assert "발급할 현금영수증이 없어요" in admin.get("/billing/receipts").text
    assert "100,000원" in admin.get("/billing/receipts?show=all").text
    assert (
        admin.patch(f"/billing/payments/{c.id}/receipt", data={"issued": "on"}).status_code == 404
    )


def test_full_pages_render(admin, db, classes):
    _student(db, "가", [classes[0]])
    svc.generate_invoices(db, svc.parse_month(None), 10)
    db.commit()
    plain = {"HX-Request": ""}
    for url in ["/billing", "/billing/receipts"]:
        r = admin.get(url, headers=plain)
        assert r.status_code == 200 and "</html>" in r.text, url
    s = db.exec(select(Invoice)).one()
    assert admin.get(f"/students/{s.student_id}", headers=plain).status_code == 200
